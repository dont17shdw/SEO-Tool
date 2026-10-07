"""Read bounded CSV/XLSX GSC exports into a deterministic, persistence-free preview.
读取大小受限的 CSV/XLSX GSC 导出文件，生成确定且不写入数据库的预览。
"""

import csv
import hashlib
import io
import re
from collections import defaultdict
from collections.abc import Iterable, Sequence
from pathlib import Path
from typing import Any
from zipfile import BadZipFile, ZipFile

from openpyxl import load_workbook

from app.imports.gsc.mapping import (
    COLUMN_ALIASES,
    FILTER_SHEET_NAMES,
    NON_PAGE_SHEET_NAMES,
    PAGE_SHEET_NAMES,
    map_columns,
    normalize_label,
)
from app.imports.gsc.schemas import ImportPreview, NormalizedGSCRow, ParsedImport, RowError
from app.normalization.gsc import (
    InvalidGSCValue,
    is_blank,
    normalize_count,
    normalize_ctr,
    normalize_position,
    normalize_url,
)

MAX_FILE_SIZE = 5 * 1024 * 1024
MAX_ROWS = 10_000
MAX_ERRORS = 100
MAX_SAMPLE_ROWS = 10
MAX_XLSX_EXPANDED_SIZE = 50 * 1024 * 1024

RawRows = list[tuple[int, Sequence[Any]]]


class ImportFileError(ValueError):
    """Expose a useful source-file error without leaking spreadsheet reader exceptions.
    公开有用的源文件错误，不泄露电子表格读取器的内部异常。
    """

    def __init__(self, code: str, message: str):
        super().__init__(message)
        self.code = code
        self.message = message


def _read_table(rows: Iterable[Sequence[Any]]) -> tuple[Sequence[Any], RawRows]:
    """Preserve source row numbers while ignoring blank lines and bounding data rows.
    忽略空白行并限制数据行数，同时保留来源行号。
    """
    headers: Sequence[Any] | None = None
    data: RawRows = []
    for row_number, row in enumerate(rows, start=1):
        if all(is_blank(cell) for cell in row):
            continue
        if headers is None:
            headers = row
        else:
            if len(data) >= MAX_ROWS:
                raise ImportFileError(
                    "too_many_rows", f"A file may contain at most {MAX_ROWS} rows."
                )
            data.append((row_number, row))
    if headers is None:
        raise ImportFileError("empty_file", "The file or Pages worksheet is empty.")
    return headers, data


def _required_mapping(headers: Sequence[Any]) -> tuple[dict[str, int], dict[str, str]]:
    """Require identifiable source columns even when their row metric values are optional.
    要求来源列可被识别，即使各行的指标值允许为空。
    """
    try:
        indices, labels = map_columns(headers)
    except ValueError as exc:
        raise ImportFileError("ambiguous_columns", str(exc)) from exc
    if "url" not in indices:
        raise ImportFileError(
            "missing_url_column",
            "A GSC Pages URL column is required (Page, URL, Top pages, or 排名靠前的网页).",
        )
    missing = COLUMN_ALIASES.keys() - indices.keys()
    if missing:
        raise ImportFileError(
            "missing_performance_columns",
            f"Cannot identify required GSC performance columns: {', '.join(sorted(missing))}.",
        )
    return indices, labels


def _read_csv(content: bytes) -> tuple[Sequence[Any], RawRows, None]:
    """Read UTF-8 CSV strictly and support the common comma, semicolon and tab separators.
    严格读取 UTF-8 CSV，并支持常见的逗号、分号及制表符分隔符。
    """
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise ImportFileError("malformed_csv", "CSV must use UTF-8 encoding.") from exc
    if "\x00" in text:
        raise ImportFileError("malformed_csv", "CSV contains invalid NUL characters.")
    try:
        try:
            dialect = csv.Sniffer().sniff(text[:8192], delimiters=",;\t")
        except csv.Error:
            dialect = csv.excel
        headers, rows = _read_table(csv.reader(io.StringIO(text, newline=""), dialect, strict=True))
    except csv.Error as exc:
        raise ImportFileError(
            "malformed_csv", "CSV cannot be read; check its delimiters and quotes."
        ) from exc
    return headers, rows, None


def _sheet_rows(sheet: Any) -> tuple[Sequence[Any], RawRows]:
    """Read worksheet records independently of potentially incorrect declared dimensions.
    读取工作表记录，不依赖可能错误的声明范围。
    """
    # Ignore unreliable worksheet dimensions; count actual records under the row limit.
    # 不依赖可能不准确的工作表范围；在行数限制下统计实际记录。
    sheet.reset_dimensions()
    return _read_table(sheet.iter_rows(values_only=True))


def _validate_window(workbook: Any) -> None:
    """Reject conflicting date metadata; absent metadata uses the documented 28-day assumption.
    拒绝冲突的日期元数据；缺失元数据时采用文档约定的 28 天假设。
    """
    for name in workbook.sheetnames:
        if normalize_label(name) not in FILTER_SHEET_NAMES:
            continue
        sheet = workbook[name]
        sheet.reset_dimensions()
        for count, row in enumerate(sheet.iter_rows(values_only=True), start=1):
            if count > MAX_ROWS:
                raise ImportFileError("too_many_rows", "The Filters worksheet is too large.")
            if not row or normalize_label(row[0]) not in {"date", "dates", "date range", "日期"}:
                continue
            value = row[1] if len(row) > 1 else None
            if is_blank(value):
                continue
            compact = re.sub(r"[\s_-]+", "", normalize_label(value))
            if compact not in {"last28days", "past28days", "28days", "过去28天", "最近28天"}:
                raise ImportFileError(
                    "unsupported_reporting_window",
                    "This import supports only the latest 28 days; "
                    "export GSC using that date filter.",
                )


def _has_page_url(rows: RawRows, url_index: int) -> bool:
    """Require actual URL evidence rather than selecting a sheet from metric headers alone.
    要求实际的 URL 证据，避免仅根据指标列名选择工作表。
    """
    for _, cells in rows:
        if url_index < len(cells):
            try:
                normalize_url(cells[url_index])
            except InvalidGSCValue:
                continue
            return True
    return False


def _read_xlsx(content: bytes) -> tuple[Sequence[Any], RawRows, str]:
    """Prefer recognized Pages sheets and require URL evidence for an unnamed fallback.
    优先选择已识别的网页工作表，并要求未命名回退表提供 URL 证据。
    """
    workbook = None
    try:
        with ZipFile(io.BytesIO(content)) as archive:
            if sum(member.file_size for member in archive.infolist()) > MAX_XLSX_EXPANDED_SIZE:
                raise ImportFileError("file_too_large", "The expanded XLSX file exceeds 50 MiB.")
        # GSC exports contain observed values; formulas are rejected rather than evaluated.
        # GSC 导出包含观测值；拒绝公式，不对公式求值。
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=False)
        _validate_window(workbook)
        for preferred in PAGE_SHEET_NAMES:
            for name in workbook.sheetnames:
                if normalize_label(name) == preferred:
                    headers, rows = _sheet_rows(workbook[name])
                    return headers, rows, name

        candidates: list[tuple[Sequence[Any], RawRows, str]] = []
        for name in workbook.sheetnames:
            if normalize_label(name) in NON_PAGE_SHEET_NAMES:
                continue
            try:
                headers, rows = _sheet_rows(workbook[name])
                indices, _ = _required_mapping(headers)
            except ImportFileError as exc:
                if exc.code == "too_many_rows":
                    raise
                continue
            if _has_page_url(rows, indices["url"]):
                candidates.append((headers, rows, name))
        if len(candidates) > 1:
            raise ImportFileError(
                "ambiguous_pages_sheet",
                "Multiple worksheets match GSC Pages; "
                "name the intended worksheet 'Pages' or '网页'.",
            )
        if candidates:
            return candidates[0]
        raise ImportFileError(
            "pages_sheet_missing",
            "No GSC Pages worksheet was found. "
            "Use Pages/网页 or matching page headers and HTTP(S) URLs.",
        )
    except ImportFileError:
        raise
    except BadZipFile as exc:
        raise ImportFileError("malformed_xlsx", "XLSX is not a readable Excel workbook.") from exc
    except Exception as exc:
        # Reader exceptions depend on malformed ZIP/XML internals and stay behind this boundary.
        # 读取器异常依赖损坏的 ZIP/XML 内部结构，应保留在该边界内部。
        raise ImportFileError("malformed_xlsx", "XLSX is not a readable Excel workbook.") from exc
    finally:
        if workbook is not None:
            workbook.close()


def _normalize_rows(
    rows: RawRows, indices: dict[str, int], header_count: int
) -> tuple[list[NormalizedGSCRow], list[RowError], set[int], set[int]]:
    """Validate every cell and block all occurrences of an exact, whitespace-trimmed URL duplicate.
    校验每个单元格，并阻止去除首尾空白后完全相同的 URL 的全部重复行。
    """
    normalized: dict[int, NormalizedGSCRow] = {}
    errors: list[RowError] = []
    invalid: set[int] = set()
    urls: dict[str, list[int]] = defaultdict(list)
    for row_number, cells in rows:
        values: dict[str, Any] = {}
        if len(cells) > header_count and any(not is_blank(cell) for cell in cells[header_count:]):
            invalid.add(row_number)
            errors.append(
                RowError(
                    row=row_number,
                    field="row",
                    code="unexpected_columns",
                    message="This row has more values than the header; check CSV quoting.",
                )
            )
        for field, index in indices.items():
            raw = cells[index] if index < len(cells) else None
            try:
                if field == "url":
                    values["url"] = normalize_url(raw)
                    urls[values["url"]].append(row_number)
                elif field in {"clicks", "impressions"}:
                    values[f"{field}_28d"] = normalize_count(raw, field)
                elif field == "ctr":
                    values["ctr"] = normalize_ctr(raw)
                else:
                    values["average_position"] = normalize_position(raw)
            except InvalidGSCValue as exc:
                invalid.add(row_number)
                errors.append(
                    RowError(row=row_number, field=field, code=exc.code, message=exc.message)
                )
        if row_number not in invalid:
            normalized[row_number] = NormalizedGSCRow(row=row_number, **values)

    duplicates: set[int] = set()
    for occurrences in urls.values():
        if len(occurrences) > 1:
            duplicates.update(occurrences)
            for row_number in occurrences:
                errors.append(
                    RowError(
                        row=row_number,
                        field="url",
                        code="duplicate_url",
                        message="This URL occurs more than once. "
                        "Remove duplicate rows before importing.",
                    )
                )
    accepted = [row for number, row in normalized.items() if number not in duplicates]
    errors.sort(key=lambda error: error.row)
    return accepted, errors, invalid - duplicates, duplicates


def parse_gsc_pages(content: bytes, filename: str) -> ParsedImport:
    """Parse, validate and normalize a source upload without writing files or database records.
    解析、校验并标准化上传的源数据，不写入文件或数据库记录。

        Duplicate and invalid rows block the entire apply; missing metric values remain NULL.
        重复行与无效行会阻止整个导入确认；缺失的指标值保持 NULL。
    """
    if len(content) > MAX_FILE_SIZE:
        raise ImportFileError("file_too_large", "Uploads must not exceed 5 MiB.")
    extension = Path(filename).suffix.lower()
    if extension not in {".csv", ".xlsx"}:
        raise ImportFileError("unsupported_file_type", "Choose a GSC Pages CSV or XLSX export.")
    if not content or not content.strip():
        raise ImportFileError("empty_file", "The uploaded file is empty.")
    headers, rows, sheet = _read_csv(content) if extension == ".csv" else _read_xlsx(content)
    indices, labels = _required_mapping(headers)
    if not rows:
        raise ImportFileError("empty_file", "The file contains headers but no page rows.")
    accepted, errors, invalid, duplicates = _normalize_rows(rows, indices, len(headers))
    preview = ImportPreview(
        detected_sheet=sheet,
        total_rows=len(rows),
        valid_rows=len(accepted),
        invalid_rows=len(invalid),
        duplicate_rows=len(duplicates),
        column_mapping=labels,
        errors=errors[:MAX_ERRORS],
        sample_rows=accepted[:MAX_SAMPLE_ROWS],
        can_apply=not invalid and not duplicates,
        preview_hash=hashlib.sha256(content).hexdigest(),
    )
    return ParsedImport(preview=preview, rows=accepted)
