"""Exercise real-format synthetic XLSX imports through existing evidence and signal APIs.
通过现有证据及信号 API 验证真实格式的合成 XLSX 导入。
"""

import io
import json
from datetime import date, timedelta

import pytest
from fastapi.testclient import TestClient
from gsc_helpers import gsc_engine as postgres_import_engine  # noqa: F401
from openpyxl import Workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.db.base import Base
from app.db.session import get_session
from app.main import create_app
from app.models import SEOOpportunity, WebsitePage

PREVIEW = "/api/v1/imports/gsc/pages/preview"
APPLY = "/api/v1/imports/gsc/pages/apply"
KNOWN = {"property_id": "sc-domain:example.com", "search_type": "web", "filters": []}
START = date(2026, 1, 1)
PREVIOUS = [
    ["https://example.com/traffic", 20, 500, "4%", 5],
    ["https://example.com/ctr", 150, 1000, "15%", 5],
    ["https://example.com/ranking", 100, 1000, "10%", 5],
    ["https://example.com/growth", 100, 1000, "10%", 5],
    ["https://example.com/no-signals", 100, 1000, "10%", 5],
]
CURRENT = [
    ["https://example.com/traffic", 8, 500, "1.6%", 5],
    ["https://example.com/ctr", 140, 1000, "14%", 5],
    ["https://example.com/ranking", 100, 1000, "10%", 9],
    ["https://example.com/growth", 105, 1500, "7%", 5],
    ["https://example.com/no-signals", 100, 1000, "10%", 5],
]


def custom_workbook(
    start, rows=PREVIOUS, *, language="en", days=range(28), device=None, label=None
):
    """Generate localized GSC Pages, Chart, and Filters without real-site data.
    生成本地化 GSC 网页、图表及筛选表，不使用真实站点数据。
    """
    chinese = language == "zh"
    workbook = Workbook()
    pages = workbook.active
    pages.title = "网页" if chinese else "Pages"
    pages.append(
        ["排名靠前的网页", "点击次数", "展示次数", "点击率", "平均排名"]
        if chinese
        else ["Top pages", "Clicks", "Impressions", "CTR", "Position"]
    )
    for row in rows:
        pages.append(row)
    chart = workbook.create_sheet("图表" if chinese else "Chart")
    chart.append(
        ["日期", "点击次数", "展示次数", "点击率", "平均排名"]
        if chinese
        else ["Date", "Clicks", "Impressions", "CTR", "Position"]
    )
    for day in days:
        chart.append([(start + timedelta(days=day)).isoformat(), 5, 50, "10%", 5])
    end = start + timedelta(days=27)
    if label is None:
        label = (
            f"自定义：{start.year}年{start.month}月{start.day}日至"
            f"{end.year}年{end.month}月{end.day}日"
            if chinese
            else f"{start.isoformat()} - {end.isoformat()}"
        )
    filters = workbook.create_sheet("过滤器" if chinese else "Filters")
    filters.append(["筛选器", "值"] if chinese else ["Filter", "Value"])
    filters.append(["搜索类型", "网络"] if chinese else ["Search type", "Web"])
    filters.append(["日期" if chinese else "Date", label])
    if device is not None:
        filters.append(["设备" if chinese else "Device", device])
    buffer = io.BytesIO()
    workbook.save(buffer)
    workbook.close()
    return buffer.getvalue()


@pytest.fixture
def workflow_client(gsc_engine, isolated_settings):
    application = create_app()

    def session():
        with Session(gsc_engine) as current:
            yield current

    application.dependency_overrides[get_session] = session
    with TestClient(application) as client:
        yield client


def import_report(client, content, *, scope=KNOWN, filename="synthetic.xlsx"):
    """Use the public preview hash and confirmation contract for every import.
    每次导入均使用公开预览哈希与确认契约。
    """
    files = {
        "file": (
            filename,
            content,
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
    }
    declaration = {} if scope is None else {"scope": json.dumps(scope)}
    preview = client.post(PREVIEW, files=files, data=declaration)
    assert preview.status_code == 200, preview.text
    applied = client.post(
        APPLY,
        files=files,
        data={**declaration, "confirmed": "true", "preview_hash": preview.json()["preview_hash"]},
    )
    assert applied.status_code == 200, applied.text
    return preview.json(), applied.json()


def stored_state(engine):
    with engine.connect() as connection:
        return {
            table.name: sorted(connection.execute(select(table)).all(), key=repr)
            for table in Base.metadata.sorted_tables
        }


@pytest.mark.parametrize("language", ["en", "zh"])
@pytest.mark.parametrize("bare_custom", [False, True])
def test_two_custom_28_day_reports_reach_ready_and_all_existing_signals(
    workflow_client, gsc_engine, language, bare_custom
):
    label = ("自定义" if language == "zh" else "Custom") if bare_custom else None
    first, first_result = import_report(
        workflow_client, custom_workbook(START, language=language, label=label)
    )
    assert first["reporting_window"] == "latest_28_days"
    assert first["coverage_status"] == "complete"
    assert first["report_scope"]["status"] == "known"
    pages = workflow_client.get("/api/v1/pages").json()["items"]
    assert len(pages) == 5
    for page in pages:
        result = workflow_client.get(f"/api/v1/pages/{page['id']}/opportunities").json()
        assert result["eligible"] is False
        assert result["candidates"] == []

    second_content = custom_workbook(
        START + timedelta(days=28), CURRENT, language=language, label=label
    )
    second, second_result = import_report(workflow_client, second_content)
    assert second["reporting_window"] == first["reporting_window"]
    assert second["observed_date_count"] == 28
    assert second["coverage_status"] == "complete"
    assert second["report_scope"]["fingerprint"] == first["report_scope"]["fingerprint"]
    assert first_result["import_run_id"] != second_result["import_run_id"]

    with Session(gsc_engine) as session:
        page = session.get(WebsitePage, pages[0]["id"])
        session.add(
            SEOOpportunity(
                page_id=page.id,
                opportunity_type="preexisting-synthetic",
                recommended_action="Synthetic existing record",
                reason="Verify that runtime analysis leaves existing records untouched",
            )
        )
        session.commit()
    baseline = stored_state(gsc_engine)
    candidates = []
    for page in pages:
        endpoint = f"/api/v1/pages/{page['id']}"
        performance = workflow_client.get(f"{endpoint}/performance?page_size=1").json()
        result = workflow_client.get(f"{endpoint}/opportunities").json()
        assert result["eligible"] is True
        assert result["evidence_readiness"] == "ready"
        assert performance["quality"]["readiness"] == "ready"
        comparison = performance["comparison"]
        assert comparison["periods_overlap"] is False
        assert comparison["scope_compatibility"] == "compatible"
        assert comparison["previous_period_start"] == "2026-01-01"
        assert comparison["previous_period_end"] == "2026-01-28"
        assert comparison["current_period_start"] == "2026-01-29"
        assert comparison["current_period_end"] == "2026-02-25"
        assert result["comparison"] == comparison
        assert performance["total"] == 2
        provenance = workflow_client.get(f"{endpoint}/provenance").json()
        assert provenance["known_provenance_count"] == 4
        assert {metric["import_run_id"] for metric in provenance["metrics"]} == {
            second_result["import_run_id"]
        }
        candidates.extend(result["candidates"])
        if page["url"].endswith("/no-signals"):
            assert result["candidates"] == []
    assert {item["opportunity_type"] for item in candidates} == {
        "traffic_decline",
        "ctr_opportunity",
        "ranking_decline",
        "impression_growth_gap",
    }
    global_result = workflow_client.get("/api/v1/opportunities").json()
    assert global_result["total"] == len(candidates) == 6
    assert global_result["items"] == sorted(
        candidates, key=lambda item: (item["url"], item["opportunity_type"], item["page_id"])
    )
    assert stored_state(gsc_engine) == baseline

    _, retry = import_report(workflow_client, second_content, filename="renamed-synthetic.xlsx")
    assert retry["already_processed"] is True
    assert retry["import_run_id"] == second_result["import_run_id"]
    assert stored_state(gsc_engine) == baseline


@pytest.mark.parametrize("language", ["en", "zh"])
@pytest.mark.parametrize(
    ("days", "coverage", "reason"),
    [([0, 9, 27], "partial", "incomplete_date_coverage"), ([], "unknown", "insufficient_history")],
)
def test_custom_ranges_never_promote_incomplete_or_unknown_dates(
    workflow_client, language, days, coverage, reason
):
    import_report(workflow_client, custom_workbook(START, language=language))
    preview, _ = import_report(
        workflow_client,
        custom_workbook(START + timedelta(days=28), CURRENT, language=language, days=days),
    )
    assert preview["coverage_status"] == coverage
    assert preview["observed_date_count"] == (len(days) if days else None)
    for page in workflow_client.get("/api/v1/pages").json()["items"]:
        result = workflow_client.get(f"/api/v1/pages/{page['id']}/opportunities").json()
        assert result["eligible"] is False
        assert result["candidates"] == []
        assert reason in result["gate_reasons"]


@pytest.mark.parametrize("language", ["en", "zh"])
def test_custom_ranges_with_unknown_scope_do_not_produce_signals(workflow_client, language):
    for start, rows in ((START, PREVIOUS), (START + timedelta(days=28), CURRENT)):
        import_report(workflow_client, custom_workbook(start, rows, language=language), scope=None)
    for page in workflow_client.get("/api/v1/pages").json()["items"]:
        result = workflow_client.get(f"/api/v1/pages/{page['id']}/opportunities").json()
        assert result["evidence_readiness"] == "limited"
        assert "unknown_report_scope" in result["gate_reasons"]
        assert result["eligible"] is False
        assert result["candidates"] == []


@pytest.mark.parametrize("language", ["en", "zh"])
def test_custom_conflicting_scope_cannot_form_ready_comparison(workflow_client, language):
    for start, rows, device in (
        (START, PREVIOUS, "mobile"),
        (START + timedelta(days=28), CURRENT, "desktop"),
    ):
        declaration = {
            **KNOWN,
            "filters": [{"dimension": "device", "operator": "equals", "value": device}],
        }
        import_report(
            workflow_client,
            custom_workbook(start, rows, language=language, device=device),
            scope=declaration,
        )
    for page in workflow_client.get("/api/v1/pages").json()["items"]:
        result = workflow_client.get(f"/api/v1/pages/{page['id']}/opportunities").json()
        quality = workflow_client.get(f"/api/v1/pages/{page['id']}/quality").json()
        assert result["comparison"] is None
        assert result["eligible"] is False
        assert result["candidates"] == []
        assert any(item["code"] == "incompatible_report_scope" for item in quality["observations"])


@pytest.mark.parametrize("language", ["en", "zh"])
def test_complete_custom_ranges_still_refuse_overlapping_comparisons(workflow_client, language):
    import_report(workflow_client, custom_workbook(START, language=language))
    import_report(
        workflow_client, custom_workbook(START + timedelta(days=7), CURRENT, language=language)
    )
    for page in workflow_client.get("/api/v1/pages").json()["items"]:
        result = workflow_client.get(f"/api/v1/pages/{page['id']}/opportunities").json()
        assert result["comparison"]["periods_overlap"] is True
        assert result["eligible"] is False
        assert result["candidates"] == []
        assert "overlapping_comparison_periods" in result["gate_reasons"]


def test_invalid_custom_range_rejected_before_database_resolution(isolated_settings):
    application = create_app()

    def unavailable():
        raise AssertionError("Invalid custom date labels must fail before any database access")

    application.dependency_overrides[get_session] = unavailable
    content = custom_workbook(START, label="2026-01-01 to 2026-01-29")
    with TestClient(application) as client:
        response = client.post(PREVIEW, files={"file": ("synthetic.xlsx", content)})
        assert response.status_code == 422
        applied = client.post(
            APPLY,
            files={"file": ("synthetic.xlsx", content)},
            data={"confirmed": "true", "preview_hash": "0" * 64},
        )
    assert applied.status_code == 422
    assert response.json()["detail"]["code"] == "unsupported_reporting_window"
    assert applied.json()["detail"]["code"] == "unsupported_reporting_window"
