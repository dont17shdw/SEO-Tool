"""Exercise scope-bound confirmation and read-only report evidence using synthetic files.
使用合成文件验证范围绑定确认及只读报告证据。
"""

import io
import json
from datetime import date, timedelta
from uuid import uuid4

import pytest
from fastapi.testclient import TestClient
from gsc_helpers import csv_file
from gsc_helpers import gsc_engine as postgres_import_engine  # noqa: F401
from openpyxl import Workbook
from sqlalchemy import event, select
from sqlalchemy.exc import SQLAlchemyError
from sqlalchemy.orm import Session

from app.db.base import Base
from app.db.session import get_session
from app.main import create_app
from app.models import ImportRun, SEOOpportunity, WebsitePage

PREVIEW = "/api/v1/imports/gsc/pages/preview"
APPLY = "/api/v1/imports/gsc/pages/apply"
KNOWN = {"property_id": "sc-domain:example.com", "search_type": "web", "filters": []}


def workbook(start, clicks=10, days=range(28), device=None):
    result = Workbook()
    pages = result.active
    pages.title = "Pages"
    pages.append(["Page", "Clicks", "Impressions", "CTR", "Position"])
    pages.append(["https://example.com/page", clicks, 100, "10%", 5])
    dates = result.create_sheet("Dates")
    dates.append(["Date"])
    for day in days:
        dates.append([start + timedelta(days=day)])
    filters = result.create_sheet("Filters")
    filters.append(["Filter", "Value"])
    filters.append(["Search type", "Web"])
    filters.append(["Date", "Last 28 days"])
    if device is not None:
        filters.append(["Device", device])
    content = io.BytesIO()
    result.save(content)
    result.close()
    return content.getvalue()


def upload(content, filename="synthetic.xlsx"):
    return {"file": (filename, content, "application/octet-stream")}


@pytest.fixture
def scope_client(gsc_engine, isolated_settings):
    application = create_app()

    def session():
        with Session(gsc_engine) as current:
            yield current

    application.dependency_overrides[get_session] = session
    with TestClient(application) as client:
        yield client


def apply(client, content, scope=KNOWN, filename="synthetic.xlsx"):
    form = {} if scope is None else {"scope": json.dumps(scope)}
    preview = client.post(PREVIEW, files=upload(content, filename), data=form)
    assert preview.status_code == 200, preview.text
    result = client.post(
        APPLY,
        files=upload(content, filename),
        data={**form, "confirmed": "true", "preview_hash": preview.json()["preview_hash"]},
    )
    assert result.status_code == 200, result.text
    return preview.json(), result.json()


@pytest.mark.parametrize("changed", ["property", "search", "filters", "completeness"])
def test_changed_scope_is_rejected_before_resolving_database(isolated_settings, changed):
    application = create_app()

    def unavailable():
        raise AssertionError("A changed scope must be rejected before database access")

    application.dependency_overrides[get_session] = unavailable
    content = csv_file(["https://example.com/page", 10, 100, "10%", 5])
    revised = dict(KNOWN)
    if changed == "property":
        revised["property_id"] = "sc-domain:other.example"
    elif changed == "search":
        revised["search_type"] = "image"
    elif changed == "filters":
        revised["filters"] = [{"dimension": "device", "operator": "equals", "value": "mobile"}]
    else:
        revised["filters"] = None
    with TestClient(application) as client:
        preview = client.post(
            PREVIEW, files=upload(content, "pages.csv"), data={"scope": json.dumps(KNOWN)}
        )
        assert preview.status_code == 200
        response = client.post(
            APPLY,
            files=upload(content, "pages.csv"),
            data={
                "scope": json.dumps(revised),
                "confirmed": "true",
                "preview_hash": preview.json()["preview_hash"],
            },
        )
    assert response.status_code == 409
    assert response.json()["detail"]["code"] == "preview_changed"


def test_evidence_origin_change_also_requires_repreview(isolated_settings):
    content = workbook(date(2026, 1, 1))
    first = {"property_id": "sc-domain:example.com", "filters": []}
    with TestClient(create_app()) as client:
        preview = client.post(PREVIEW, files=upload(content), data={"scope": json.dumps(first)})
        assert preview.status_code == 200
        other = client.post(PREVIEW, files=upload(content), data={"scope": json.dumps(KNOWN)})
        assert other.status_code == 200
        assert (
            preview.json()["report_scope"]["fingerprint"]
            == other.json()["report_scope"]["fingerprint"]
        )
        assert preview.json()["preview_hash"] != other.json()["preview_hash"]
        changed = client.post(
            APPLY,
            files=upload(content),
            data={
                "scope": json.dumps(KNOWN),
                "confirmed": "true",
                "preview_hash": preview.json()["preview_hash"],
            },
        )
        assert changed.status_code == 409


@pytest.mark.parametrize(
    "scope", ["broken json", "null", "[]", '{"unexpected":true}', '{"search_type":"unsupported"}']
)
def test_invalid_scope_declarations_are_public_errors(isolated_settings, scope):
    with TestClient(create_app()) as client:
        response = client.post(
            PREVIEW,
            files=upload(csv_file(["https://example.com/page", 1, 10, "10%", 5]), "pages.csv"),
            data={"scope": scope},
        )
    assert response.status_code == 422
    assert response.json()["detail"]["code"] == "invalid_report_scope"


def test_observed_conflict_blocks_import_without_database(isolated_settings):
    declaration = {**KNOWN, "search_type": "image"}
    with TestClient(create_app()) as client:
        response = client.post(
            PREVIEW,
            files=upload(workbook(date(2026, 1, 1))),
            data={"scope": json.dumps(declaration)},
        )
    assert response.status_code == 422
    assert response.json()["detail"]["code"]


def stored_rows(engine):
    """Capture every table column so read-only assertions include timestamps and links.
    捕获每个表字段，使只读断言涵盖时间戳及关联。
    """
    with engine.connect() as connection:
        return {
            table.name: sorted(connection.execute(select(table)).all(), key=repr)
            for table in Base.metadata.sorted_tables
        }


def test_complete_scope_and_coverage_reach_ready_without_extra_queries(scope_client, gsc_engine):
    previous, first_run = apply(scope_client, workbook(date(2026, 1, 1), clicks=10))
    current, second_run = apply(scope_client, workbook(date(2026, 2, 1), clicks=20))
    assert current["coverage_status"] == previous["coverage_status"] == "complete"
    assert current["observed_date_count"] == 28 and current["dates_consecutive"] is True
    pages = scope_client.get("/api/v1/pages").json()["items"]
    page_id = pages[0]["id"]
    baseline = stored_rows(gsc_engine)
    queries = []

    def capture(connection, cursor, statement, parameters, context, executemany):
        if statement.lstrip().upper().startswith("SELECT"):
            queries.append(statement)

    event.listen(gsc_engine, "before_cursor_execute", capture)
    try:
        performance = scope_client.get(f"/api/v1/pages/{page_id}/performance?page_size=1").json()
        assert len(queries) == 3
        queries.clear()
        quality = scope_client.get(f"/api/v1/pages/{page_id}/quality").json()
        assert len(queries) == 2
        details = scope_client.get(f"/api/v1/imports/{second_run['import_run_id']}").json()
        scope_client.get(f"/api/v1/imports/{first_run['import_run_id']}/quality").raise_for_status()
        provenance = scope_client.get(f"/api/v1/pages/{page_id}/provenance").json()
    finally:
        event.remove(gsc_engine, "before_cursor_execute", capture)
    assert performance["quality"] == quality
    assert quality["readiness"] == "ready"
    assert performance["comparison"]["scope_compatibility"] == "compatible"
    assert performance["comparison"]["previous_coverage_status"] == "complete"
    assert details["report_scope"] == current["report_scope"]
    assert details["site_id"] == pages[0]["site_id"]
    assert details["report_scope"]["property_status"] == "known"
    assert performance["items"][0]["report_scope"] == previous["report_scope"]
    assert provenance == performance["provenance"]
    assert provenance["known_provenance_count"] == 4
    assert stored_rows(gsc_engine) == baseline
    with Session(gsc_engine) as session:
        assert session.scalars(select(SEOOpportunity)).all() == []


def test_sparse_date_evidence_reduces_selected_readiness(scope_client):
    apply(scope_client, workbook(date(2026, 1, 1)))
    apply(scope_client, workbook(date(2026, 2, 1), days=[0, 9, 27]))
    page_id = scope_client.get("/api/v1/pages").json()["items"][0]["id"]
    data = scope_client.get(f"/api/v1/pages/{page_id}/performance").json()
    assert data["comparison"]["scope_compatibility"] == "compatible"
    assert data["comparison"]["current_observed_date_count"] == 3
    assert data["comparison"]["current_coverage_status"] == "partial"
    assert data["quality"]["readiness"] == "limited"
    assert "incomplete_date_coverage" in data["quality"]["readiness_reasons"]


def test_partial_scope_remains_descriptive_and_visibly_unknown(scope_client):
    first, _ = apply(scope_client, workbook(date(2026, 1, 1)), scope=None)
    assert first["report_scope"]["property_id"] is None
    assert first["report_scope"]["search_type"] == "web"
    assert first["report_scope"]["filters_complete"] is False
    apply(scope_client, workbook(date(2026, 2, 1)), scope=None)
    page_id = scope_client.get("/api/v1/pages").json()["items"][0]["id"]
    data = scope_client.get(f"/api/v1/pages/{page_id}/performance").json()
    assert data["comparison"]["scope_compatibility"] == "unknown"
    assert data["quality"]["readiness"] == "limited"
    assert "unknown_report_scope" in data["quality"]["readiness_reasons"]


def test_conflicting_scope_cannot_replace_an_earlier_compatible_pair(scope_client):
    _, first = apply(scope_client, workbook(date(2026, 1, 1)))
    _, second = apply(scope_client, workbook(date(2026, 2, 1)))
    mobile = {
        **KNOWN,
        "filters": [{"dimension": "device", "operator": "equals", "value": "mobile"}],
    }
    _, third = apply(scope_client, workbook(date(2026, 3, 1), device="Mobile"), scope=mobile)
    page_id = scope_client.get("/api/v1/pages").json()["items"][0]["id"]
    data = scope_client.get(f"/api/v1/pages/{page_id}/performance").json()
    assert data["comparison"]["previous_period_start"] == "2026-01-01"
    assert data["comparison"]["current_period_start"] == "2026-02-01"
    assert data["quality"]["readiness"] == "ready"
    assert data["provenance"]["metrics"][0]["import_run_id"] == third["import_run_id"]
    assert first["import_run_id"] != second["import_run_id"] != third["import_run_id"]


def test_legacy_run_is_unknown_without_mutating_or_assigning_site(scope_client, gsc_engine):
    with Session(gsc_engine) as session:
        page = WebsitePage(url="https://example.com/legacy", clicks_28d=3)
        run = ImportRun(
            file_hash="a" * 64,
            filename="example.com.xlsx",
            total_rows=1,
            skipped_count=1,
            period_start=date(2026, 1, 1),
            period_end=date(2026, 1, 28),
        )
        session.add_all([page, run])
        session.commit()
        run_id = run.id
    baseline = stored_rows(gsc_engine)
    data = scope_client.get(f"/api/v1/imports/{run_id}").json()
    assert data["site_id"] is None
    assert data["report_scope"]["status"] == "unknown"
    assert data["report_scope"]["property_id"] is None
    assert data["report_scope"]["filters_complete"] is False
    assert data["coverage_status"] == "unknown"
    assert data["observed_date_count"] is data["dates_consecutive"] is None
    assert stored_rows(gsc_engine) == baseline


def test_import_details_errors_are_read_only_and_generic(scope_client, monkeypatch):
    assert scope_client.get(f"/api/v1/imports/{uuid4()}").status_code == 404
    assert scope_client.get("/api/v1/imports/not-a-uuid").status_code == 422

    def unavailable(*args, **kwargs):
        raise SQLAlchemyError("private database detail")

    monkeypatch.setattr(Session, "get", unavailable)
    response = scope_client.get(f"/api/v1/imports/{uuid4()}")
    assert response.status_code == 503
    assert "private database detail" not in response.text
